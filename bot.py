"""Excel sipariş tablosundan günlük özet çıkarıp Telegram'a gönderen bot.

Kullanım:
  python bot.py --onizleme            Mesajı Telegram'a göndermeden ekrana yazar
  python bot.py --gonder              Bugünün özetini Telegram'a gönderir (zamanlanmış görev için)
  python bot.py --gonder --tarih 2026-10-14
  python bot.py --dinle               Komutları dinler: /ozet /dun /hafta /bekleyen /urunler /yardim
"""
import argparse
import json
import os
import sys
import time
import urllib.parse
import urllib.request
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

KLASOR = Path(__file__).parent
GUNLER = ["Pazartesi", "Salı", "Çarşamba", "Perşembe", "Cuma", "Cumartesi", "Pazar"]
AYLAR = ["Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran", "Temmuz", "Ağustos", "Eylül", "Ekim", "Kasım", "Aralık"]


# ---------- Ayarlar ----------
def ayarlari_yukle():
    """.env dosyasını (varsa) ortam değişkenlerine yükler."""
    env = KLASOR / ".env"
    if env.exists():
        for satir in env.read_text(encoding="utf8").splitlines():
            if "=" in satir and not satir.lstrip().startswith("#"):
                anahtar, deger = satir.split("=", 1)
                os.environ.setdefault(anahtar.strip(), deger.strip())
    return {
        "token": os.environ.get("TELEGRAM_TOKEN", ""),
        "chat_id": os.environ.get("TELEGRAM_CHAT_ID", ""),
        "excel": os.environ.get("EXCEL_DOSYASI", str(KLASOR / "siparisler.xlsx")),
    }


# ---------- Excel okuma ----------
@dataclass
class Siparis:
    zaman: datetime
    no: str
    musteri: str
    urun: str
    adet: int
    fiyat: float
    durum: str
    odeme: str

    @property
    def tutar(self):
        return self.adet * self.fiyat


def siparisleri_oku(dosya):
    ws = load_workbook(dosya, data_only=True)["Siparişler"]
    basliklar = [h.value for h in ws[1]]
    i = {ad: basliklar.index(ad) for ad in ["Tarih", "Sipariş No", "Müşteri", "Ürün", "Adet", "Birim Fiyat", "Durum", "Ödeme"]}
    siparisler = []
    for satir in ws.iter_rows(min_row=2, values_only=True):
        if not satir[i["Tarih"]]:
            continue
        siparisler.append(Siparis(
            zaman=satir[i["Tarih"]], no=satir[i["Sipariş No"]], musteri=satir[i["Müşteri"]], urun=satir[i["Ürün"]],
            adet=int(satir[i["Adet"]]), fiyat=float(satir[i["Birim Fiyat"]]), durum=satir[i["Durum"]], odeme=satir[i["Ödeme"]],
        ))
    return siparisler


# ---------- Rapor ----------
def tl(sayi):
    return f"{sayi:,.0f} ₺".replace(",", ".")


def tarih_yaz(g):
    return f"{g.day} {AYLAR[g.month - 1]} {GUNLER[g.weekday()]}"


def ozet(siparisler, baslangic, bitis=None):
    """baslangic–bitis (dahil) arasındaki siparişlerin özet mesajı."""
    bitis = bitis or baslangic
    aralik = [s for s in siparisler if baslangic <= s.zaman.date() <= bitis]
    gecerli = [s for s in aralik if s.durum != "İptal"]
    iptal = [s for s in aralik if s.durum == "İptal"]
    baslik = tarih_yaz(baslangic) if baslangic == bitis else f"{baslangic:%d.%m} – {bitis:%d.%m.%Y}"

    if not aralik:
        return f"📊 *{baslik}*\n\nBu tarihte sipariş yok."

    ciro = sum(s.tutar for s in gecerli)
    satir = [f"📊 *Satış Özeti — {baslik}*", ""]
    satir.append(f"🧾 Sipariş: *{len(gecerli)}*" + (f"  (iptal: {len(iptal)})" if iptal else ""))
    satir.append(f"💰 Ciro: *{tl(ciro)}*")
    if gecerli:
        satir.append(f"🛒 Ortalama sepet: {tl(ciro / len(gecerli))}")

    # Önceki aynı uzunluktaki dönemle karşılaştır. Gün bitmediyse önceki dönemi de aynı saatte kes.
    gun = (bitis - baslangic).days + 1
    simdi = datetime.now()
    devam_ediyor = bitis >= simdi.date()
    once = [s for s in siparisler
            if baslangic - timedelta(days=gun) <= s.zaman.date() < baslangic and s.durum != "İptal"
            and not (devam_ediyor and s.zaman.date() == bitis - timedelta(days=gun) and s.zaman.time() > simdi.time())]
    onceki_ciro = sum(s.tutar for s in once)
    if onceki_ciro:
        fark = (ciro - onceki_ciro) / onceki_ciro * 100
        etiket = ("Dün bu saate göre" if gun == 1 else "Önceki döneme göre (aynı saate kadar)") if devam_ediyor else (
            "Önceki güne göre" if gun == 1 else f"Önceki {gun} güne göre")
        isaret = "+" if fark >= 0 else "−"
        satir.append(f"{'📈' if fark >= 0 else '📉'} {etiket}: {isaret}%{abs(fark):.0f}")

    adetler = Counter()
    for s in gecerli:
        adetler[s.urun] += s.adet
    if adetler:
        satir += ["", "🏆 *En çok satanlar*"]
        for sira, (urun, adet) in enumerate(adetler.most_common(3), 1):
            satir.append(f"{sira}. {urun} — {adet} adet")

    durumlar = Counter(s.durum for s in aralik)
    satir += ["", "📦 *Durum*"] + [f"• {d}: {n}" for d, n in durumlar.most_common()]
    return "\n".join(satir)


def bekleyenler(siparisler):
    bekleyen = sorted((s for s in siparisler if s.durum == "Hazırlanıyor"), key=lambda s: s.zaman)
    if not bekleyen:
        return "✅ Hazırlanmayı bekleyen sipariş yok."
    satir = [f"⏳ *Hazırlanmayı bekleyen {len(bekleyen)} sipariş*", ""]
    for s in bekleyen[:15]:
        satir.append(f"• `{s.no}` {s.zaman:%d.%m %H:%M} — {s.musteri}, {s.urun} ×{s.adet}")
    if len(bekleyen) > 15:
        satir.append(f"… ve {len(bekleyen) - 15} sipariş daha")
    return "\n".join(satir)


def urun_raporu(siparisler, gun=7):
    sinir = date.today() - timedelta(days=gun - 1)
    ciro, adet = defaultdict(float), Counter()
    for s in siparisler:
        if s.zaman.date() >= sinir and s.durum != "İptal":
            ciro[s.urun] += s.tutar
            adet[s.urun] += s.adet
    satir = [f"📦 *Son {gun} gün — ürün bazında*", ""]
    for urun, c in sorted(ciro.items(), key=lambda x: -x[1]):
        satir.append(f"• {urun}: {adet[urun]} adet, {tl(c)}")
    return "\n".join(satir)


YARDIM = (
    "🤖 *Sipariş Özet Botu*\n\n"
    "/ozet — bugünün özeti\n"
    "/dun — dünün özeti\n"
    "/hafta — son 7 günün özeti\n"
    "/bekleyen — hazırlanmayı bekleyen siparişler\n"
    "/urunler — son 7 gün ürün bazında satış\n"
    "/yardim — bu mesaj"
)


def komut_cevabi(metin, siparisler):
    komut = metin.strip().split()[0].split("@")[0].lower() if metin.strip() else ""
    bugun = date.today()
    if komut in ("/ozet", "/start"):
        return ozet(siparisler, bugun) if komut == "/ozet" else YARDIM
    if komut == "/dun":
        return ozet(siparisler, bugun - timedelta(days=1))
    if komut == "/hafta":
        return ozet(siparisler, bugun - timedelta(days=6), bugun)
    if komut == "/bekleyen":
        return bekleyenler(siparisler)
    if komut == "/urunler":
        return urun_raporu(siparisler)
    return YARDIM


# ---------- Telegram ----------
def telegram(token, metod, **params):
    url = f"https://api.telegram.org/bot{token}/{metod}"
    veri = urllib.parse.urlencode(params).encode()
    with urllib.request.urlopen(url, data=veri, timeout=40) as r:
        cevap = json.load(r)
    if not cevap.get("ok"):
        raise RuntimeError(cevap)
    return cevap["result"]


def gonder(ayar, mesaj, chat_id=None):
    telegram(ayar["token"], "sendMessage", chat_id=chat_id or ayar["chat_id"], text=mesaj, parse_mode="Markdown")


def dinle(ayar):
    """Long polling ile komutları dinler. Sadece ayarlı chat_id'ye cevap verir."""
    print("Bot dinliyor… (durdurmak için Ctrl+C)")
    son = 0
    while True:
        try:
            for g in telegram(ayar["token"], "getUpdates", offset=son + 1, timeout=30):
                son = g["update_id"]
                msj = g.get("message") or {}
                chat = str(msj.get("chat", {}).get("id", ""))
                if not msj.get("text"):
                    continue
                if not ayar["chat_id"]:
                    print(f"Mesaj geldi, chat id: {chat}  → .env'deki TELEGRAM_CHAT_ID'ye yazın")
                elif chat != ayar["chat_id"]:
                    print(f"Yetkisiz sohbetten mesaj yok sayıldı: {chat}")
                    continue
                gonder(ayar, komut_cevabi(msj["text"], siparisleri_oku(ayar["excel"])), chat)
        except KeyboardInterrupt:
            break
        except Exception as hata:  # Ağ kesintisinde botun çökmemesi için
            print("Hata:", hata, file=sys.stderr)
            time.sleep(5)


def main():
    p = argparse.ArgumentParser(description="Excel sipariş tablosundan Telegram'a günlük özet")
    p.add_argument("--onizleme", action="store_true", help="mesajı göndermeden ekrana yaz")
    p.add_argument("--gonder", action="store_true", help="özeti Telegram'a gönder")
    p.add_argument("--dinle", action="store_true", help="komutları dinle")
    p.add_argument("--tarih", type=date.fromisoformat, default=date.today(), help="YYYY-AA-GG (varsayılan: bugün)")
    a = p.parse_args()
    ayar = ayarlari_yukle()

    if a.onizleme or not (a.gonder or a.dinle):
        siparisler = siparisleri_oku(ayar["excel"])
        for metin in (ozet(siparisler, a.tarih), bekleyenler(siparisler)):
            print(metin, end="\n\n")
        return
    if not ayar["token"]:
        sys.exit("TELEGRAM_TOKEN eksik. .env.example dosyasını .env olarak kopyalayıp doldurun.")
    if a.gonder:
        if not ayar["chat_id"]:
            sys.exit("TELEGRAM_CHAT_ID eksik. Botu --dinle ile başlatıp bota yazın, chat id konsola düşer.")
        gonder(ayar, ozet(siparisleri_oku(ayar["excel"]), a.tarih))
        print("Özet gönderildi.")
    if a.dinle:
        dinle(ayar)


if __name__ == "__main__":
    main()
