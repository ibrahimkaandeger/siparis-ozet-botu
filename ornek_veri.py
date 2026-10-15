"""Demo için son 7 günün örnek sipariş tablosunu oluşturur: siparisler.xlsx"""
import random
from datetime import date, datetime, time, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

URUNLER = {
    "Filtre Kahve 250 g": 320,
    "Espresso Çekirdek 1 kg": 1150,
    "V60 Demleme Seti": 890,
    "Kahve Değirmeni (Manuel)": 1450,
    "French Press 600 ml": 640,
    "Seramik Kupa": 280,
    "Hediye Kutusu": 1250,
}
MUSTERILER = ["Ayşe K.", "Mehmet Y.", "Zeynep A.", "Can D.", "Elif S.", "Burak T.", "Selin Ö.", "Emre B.", "Deniz Ç.", "Ece G."]
# Gün eskidikçe siparişler ilerler: bugün çoğu hazırlanıyor, 2+ gün önce çoğu teslim edilmiş
DURUMLAR = {
    0: ["Hazırlanıyor"] * 7 + ["Kargoda"] * 2 + ["İptal"],
    1: ["Hazırlanıyor"] * 2 + ["Kargoda"] * 6 + ["Teslim edildi"] * 1 + ["İptal"],
    2: ["Kargoda"] * 3 + ["Teslim edildi"] * 6 + ["İptal"],
}
VARSAYILAN_DURUM = ["Teslim edildi"] * 12 + ["İptal"]
ODEMELER = ["Kredi kartı", "Kredi kartı", "Havale/EFT", "Kapıda ödeme"]


def olustur(dosya="siparisler.xlsx", gun_sayisi=7, seed=42):
    random.seed(seed)
    wb = Workbook()
    ws = wb.active
    ws.title = "Siparişler"
    basliklar = ["Tarih", "Sipariş No", "Müşteri", "Ürün", "Adet", "Birim Fiyat", "Tutar", "Durum", "Ödeme"]
    ws.append(basliklar)

    no = 1040
    bugun = date.today()
    for g in range(gun_sayisi - 1, -1, -1):
        gun = bugun - timedelta(days=g)
        for _ in range(random.randint(8, 16)):
            no += 1
            urun = random.choice(list(URUNLER))
            adet = random.choices([1, 2, 3], weights=[70, 22, 8])[0]
            fiyat = URUNLER[urun]
            # Bugün için şu anki saatten sonraya sipariş yazma
            son_saat = min(22, datetime.now().hour) if g == 0 else 22
            saat = time(random.randint(8, max(8, son_saat)), random.randint(0, 59))
            durum = random.choice(DURUMLAR.get(g, VARSAYILAN_DURUM))
            ws.append([datetime.combine(gun, saat), f"SP-{no}", random.choice(MUSTERILER), urun, adet, fiyat,
                       f"=E{ws.max_row + 1}*F{ws.max_row + 1}", durum, random.choice(ODEMELER)])

    # Biçimlendirme
    for hucre in ws[1]:
        hucre.font = Font(bold=True, color="FFFFFF")
        hucre.fill = PatternFill("solid", fgColor="0F766E")
        hucre.alignment = Alignment(horizontal="center")
    for satir in ws.iter_rows(min_row=2):
        satir[0].number_format = "DD.MM.YYYY HH:MM"
        satir[5].number_format = '#,##0 "₺"'
        satir[6].number_format = '#,##0 "₺"'
    for kolon, genislik in zip("ABCDEFGHI", [17, 12, 13, 26, 7, 12, 12, 15, 14]):
        ws.column_dimensions[kolon].width = genislik
    tablo = Table(displayName="Siparisler", ref=f"A1:I{ws.max_row}")
    tablo.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(tablo)
    ws.freeze_panes = "A2"
    wb.save(dosya)
    return ws.max_row - 1


if __name__ == "__main__":
    print(f"siparisler.xlsx oluşturuldu: {olustur()} sipariş")
